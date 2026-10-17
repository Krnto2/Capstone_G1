"""
Exportación del reporte de cumplimiento (CU-C15, RF-10): PDF, XLSX y CSV.

El PDF se genera con WeasyPrint a partir de una plantilla HTML. WeasyPrint
necesita las bibliotecas Pango del sistema, incluidas en la imagen Docker; si
no están (por ejemplo, en un equipo Windows de desarrollo) se entrega la misma
plantilla lista para imprimir o guardar como PDF desde el navegador.
"""
import base64
import csv
import io
from functools import lru_cache
from pathlib import Path

from django.conf import settings
from django.contrib.staticfiles import finders

from django.template.loader import render_to_string
from django.utils import timezone

ENCABEZADOS = ["Código", "Elemento", "Categoría", "Ubicación", "Plan", "Responsable", "Fecha programada",
               "Fecha límite", "Estado", "Resultado", "Ejecutada por", "Fecha de ejecución", "En plazo"]


def _fila(revision):
    material = revision.material
    en_plazo = revision.en_plazo
    return [
        material.codigo_institucional or "",
        material.nombre,
        material.categoria.nombre,
        material.ubicacion,
        str(revision.plan),
        revision.usuario_responsable.nombre_completo,
        revision.fecha_programada.strftime("%d-%m-%Y"),
        revision.fecha_limite.strftime("%d-%m-%Y"),
        revision.get_estado_display(),
        revision.get_resultado_global_display() if revision.resultado_global else "",
        revision.usuario_ejecutor.nombre_completo if revision.usuario_ejecutor else "",
        timezone.localtime(revision.fecha_hora_termino).strftime("%d-%m-%Y %H:%M") if revision.fecha_hora_termino
        else "",
        "" if en_plazo is None else ("Sí" if en_plazo else "No"),
    ]


def nombre_archivo(filtros, extension):
    return f"cumplimiento_sirema_{filtros.periodo}.{extension}"


def exportar_csv(indicadores, revisiones):
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow(["Reporte de cumplimiento SIREMA", indicadores.filtros.etiqueta])
    escritor.writerow(["Cumplimiento", f"{indicadores.porcentaje} %" if indicadores.porcentaje is not None else "—",
                       f"{indicadores.en_plazo} de {indicadores.exigibles} en plazo"])
    escritor.writerow([])
    escritor.writerow(ENCABEZADOS)
    for revision in revisiones:
        escritor.writerow(_fila(revision))
    return ("﻿" + salida.getvalue()).encode("utf-8")


def exportar_xlsx(indicadores, revisiones):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="111214")
    blanco = Font(bold=True, color="FFFFFF")

    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(["Reporte de cumplimiento SIREMA"])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append([indicadores.filtros.etiqueta])
    resumen.append([f"Generado el {timezone.localtime():%d-%m-%Y %H:%M}"])
    resumen.append([])
    filas = [
        ("Cumplimiento del período (%)", indicadores.porcentaje if indicadores.porcentaje is not None else "—"),
        ("Revisiones en plazo", indicadores.en_plazo),
        ("Revisiones con fecha límite exigible", indicadores.exigibles),
        ("Realizadas fuera de plazo", indicadores.fuera_de_plazo),
        ("Revisiones vencidas", len(indicadores.vencidas)),
        ("Hallazgos abiertos", len(indicadores.hallazgos_abiertos)),
        ("Hallazgos de criticidad alta", indicadores.hallazgos_alta),
        ("Elementos activos sin plan", indicadores.sin_plan),
        ("Elementos activos", indicadores.total_activos),
    ]
    for etiqueta, valor in filas:
        resumen.append([etiqueta, valor])
    resumen.append([])
    resumen.append(["Cumplimiento por unidad", "En plazo", "Exigibles", "%"])
    for celda in resumen[resumen.max_row]:
        celda.font = negrita
    for fila in indicadores.por_unidad:
        resumen.append([fila.etiqueta, fila.en_plazo, fila.exigibles, fila.porcentaje])
    resumen.append([])
    resumen.append(["Cumplimiento por categoría", "En plazo", "Exigibles", "%"])
    for celda in resumen[resumen.max_row]:
        celda.font = negrita
    for fila in indicadores.por_categoria:
        resumen.append([fila.etiqueta, fila.en_plazo, fila.exigibles, fila.porcentaje])
    resumen.column_dimensions["A"].width = 42
    for letra in "BCD":
        resumen.column_dimensions[letra].width = 14

    detalle = libro.create_sheet("Revisiones")
    detalle.append(ENCABEZADOS)
    for revision in revisiones:
        detalle.append(_fila(revision))

    hallazgos = libro.create_sheet("Hallazgos abiertos")
    hallazgos.append(["Código", "Elemento", "Ubicación", "Descripción", "Criticidad", "Estado", "Apertura"])
    for h in indicadores.hallazgos_abiertos:
        material = h.revision.material
        hallazgos.append([material.codigo_institucional or "", material.nombre, material.ubicacion, h.descripcion,
                          h.get_criticidad_display(), h.get_estado_display(),
                          timezone.localtime(h.fecha_apertura).strftime("%d-%m-%Y %H:%M")])

    for hoja in (detalle, hallazgos):
        for celda in hoja[1]:
            celda.font = blanco
            celda.fill = relleno
            celda.alignment = Alignment(vertical="center")
        hoja.freeze_panes = "A2"
        for i, _ in enumerate(hoja[1], start=1):
            hoja.column_dimensions[get_column_letter(i)].width = 20

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


@lru_cache(maxsize=1)
def _logo_data_uri():
    """Emblema incrustado en el HTML: sirve igual para WeasyPrint y para la vista imprimible."""
    ruta = finders.find("img/logo-8.png")
    return "data:image/png;base64," + base64.b64encode(Path(ruta).read_bytes()).decode() if ruta else ""


def _fuentes_oswald():
    """Rutas de Oswald: archivo local para WeasyPrint; la vista imprimible usa las URL estáticas."""
    rutas = {}
    for peso in (500, 700):
        ruta = finders.find(f"vendor/oswald/oswald-latin-{peso}-normal.woff2")
        rutas[peso] = Path(ruta).as_uri() if ruta else ""
    return rutas


def html_reporte(indicadores, revisiones, para_imprimir=False):
    return render_to_string("reportes/reporte_pdf.html", {
        "ind": indicadores, "revisiones": revisiones, "generado": timezone.localtime(),
        "para_imprimir": para_imprimir, "logo": _logo_data_uri(), "SIREMA_VERSION": settings.SIREMA_VERSION,
        "fuentes": None if para_imprimir else _fuentes_oswald(),
    })


def exportar_pdf(indicadores, revisiones):
    """Devuelve (bytes, True) con WeasyPrint o (html, False) si no está disponible."""
    try:
        from weasyprint import HTML
    except (ImportError, OSError):
        return html_reporte(indicadores, revisiones, para_imprimir=True), False
    return HTML(string=html_reporte(indicadores, revisiones)).write_pdf(), True
