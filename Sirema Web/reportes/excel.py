"""
Reporte de cumplimiento en XLSX con la identidad de la Octava Compañía (CU-C15).

Replica el diseño del PDF: emblema, franja negra y roja, indicadores con color,
encabezados de tabla negros y filas alternadas. Las fechas se escriben como
fechas reales de Excel para poder ordenar y filtrar, y cada hoja queda lista
para imprimir en horizontal.

Los títulos usan Arial Narrow (incluida en Office), equivalente condensado de
la Oswald de la interfaz, que no viene instalada en los computadores.
"""
import io

from django.conf import settings
from django.contrib.staticfiles import finders
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

NEGRO, AZUL, ROJO, AMBAR = "0E0F10", "1A6BB5", "D3121A", "B26C00"
GRIS_TEXTO, GRIS_BORDE, CEBRA = "4F555D", "D9DDE2", "F6F7F8"
VERDE_CONFORME, ROJO_FALLA, AMBAR_OBS = "1D7A35", "B4161C", "8A5300"
TITULOS, CUERPO = "Arial Narrow", "Calibri"

COMPANIA = "BOMBA HUACHIPATO · OCTAVA COMPAÑÍA DE BOMBEROS DE TALCAHUANO"
PIE = "SIREMA · Octava Compañía de Bomberos de Talcahuano"

relleno = lambda color: PatternFill("solid", fgColor=color)  # noqa: E731
fino = Side(style="thin", color=GRIS_BORDE)
BORDE_CELDA = Border(bottom=Side(style="thin", color="E1E4E8"))


def _fecha(valor):
    """Excel no admite zonas horarias: se escribe la hora local de Chile."""
    if valor is None:
        return None
    if hasattr(valor, "tzinfo") and valor.tzinfo is not None:
        return timezone.localtime(valor).replace(tzinfo=None)
    return valor


def _logo(hoja, alto_px=62):
    ruta = finders.find("img/logo-8-chico.png")
    if not ruta:
        return
    imagen = ImagenExcel(ruta)
    proporcion = imagen.width / imagen.height
    imagen.height, imagen.width = alto_px, round(alto_px * proporcion)
    hoja.add_image(imagen, "A1")


def _cabecera(hoja, titulo, filtros, ultima_col):
    """Emblema, compañía, título, filtros y franja negra y roja (filas 1 a 5)."""
    letra_fin = get_column_letter(ultima_col)
    _logo(hoja)
    hoja.row_dimensions[1].height = 18
    hoja.row_dimensions[2].height = 30
    hoja.row_dimensions[3].height = 18
    hoja["C1"] = COMPANIA
    hoja["C1"].font = Font(name=TITULOS, size=10, bold=True, color=AZUL)
    hoja["C2"] = titulo.upper()
    hoja["C2"].font = Font(name=TITULOS, size=20, bold=True, color=NEGRO)
    hoja["C3"] = filtros
    hoja["C3"].font = Font(name=CUERPO, size=10, color=GRIS_TEXTO)
    hoja[f"{letra_fin}1"] = f"SIREMA v{settings.SIREMA_VERSION}"
    hoja[f"{letra_fin}2"] = f"Generado el {timezone.localtime():%d-%m-%Y %H:%M}"
    for celda in (f"{letra_fin}1", f"{letra_fin}2"):
        hoja[celda].font = Font(name=CUERPO, size=9, color=GRIS_TEXTO)
        hoja[celda].alignment = Alignment(horizontal="right")
    for fila, color in ((4, NEGRO), (5, ROJO)):
        hoja.row_dimensions[fila].height = 3.5
        for col in range(1, ultima_col + 1):
            hoja.cell(row=fila, column=col).fill = relleno(color)


def _preparar_impresion(hoja, filas_titulo=None):
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.6)
    hoja.oddFooter.left.text = PIE
    hoja.oddFooter.left.size = 8
    hoja.oddFooter.right.text = "Página &P de &N"
    hoja.oddFooter.right.size = 8
    if filas_titulo:
        hoja.print_title_rows = filas_titulo
    hoja.sheet_view.showGridLines = False


def _encabezado_tabla(hoja, fila, columnas):
    """columnas: lista de (texto, columna_inicial, columna_final)."""
    hoja.row_dimensions[fila].height = 22
    for texto, desde, hasta in columnas:
        if hasta > desde:
            hoja.merge_cells(start_row=fila, start_column=desde, end_row=fila, end_column=hasta)
        for col in range(desde, hasta + 1):
            hoja.cell(row=fila, column=col).fill = relleno(NEGRO)
        celda = hoja.cell(row=fila, column=desde, value=texto.upper())
        celda.font = Font(name=TITULOS, size=10, bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center", horizontal="left", indent=1)


def _titulo_seccion(hoja, fila, texto):
    hoja.row_dimensions[fila].height = 22
    celda = hoja.cell(row=fila, column=1, value=texto.upper())
    celda.font = Font(name=TITULOS, size=12, bold=True, color=NEGRO)
    celda.alignment = Alignment(vertical="center", indent=1)
    celda.border = Border(left=Side(style="thick", color=AZUL))


# ---------------------------------------------------------------- Resumen ----

def _tarjeta(hoja, col, etiqueta, valor, detalle, color, formato=None):
    """Indicador en dos columnas (filas 7 a 10), con franja superior de color."""
    hoja.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col + 1)
    hoja.merge_cells(start_row=9, start_column=col, end_row=9, end_column=col + 1)
    hoja.merge_cells(start_row=10, start_column=col, end_row=10, end_column=col + 1)
    for c in (col, col + 1):
        hoja.cell(row=7, column=c).fill = relleno(color)
        for fila in (8, 9, 10):
            izquierda = fino if c == col else None
            derecha = fino if c == col + 1 else None
            abajo = fino if fila == 10 else None
            hoja.cell(row=fila, column=c).border = Border(left=izquierda, right=derecha, bottom=abajo)
    e = hoja.cell(row=8, column=col, value=etiqueta.upper())
    e.font = Font(name=TITULOS, size=9, bold=True, color=GRIS_TEXTO)
    e.alignment = Alignment(indent=1, vertical="bottom")
    v = hoja.cell(row=9, column=col, value=valor)
    v.font = Font(name=TITULOS, size=24, bold=True, color=NEGRO if color in (AZUL, NEGRO) else color)
    v.alignment = Alignment(indent=1, vertical="center", horizontal="left")
    if formato:
        v.number_format = formato
    d = hoja.cell(row=10, column=col, value=detalle)
    d.font = Font(name=CUERPO, size=9, color=GRIS_TEXTO)
    d.alignment = Alignment(indent=1, vertical="top")


def _tabla_cumplimiento(hoja, fila, titulo, filas):
    _titulo_seccion(hoja, fila, titulo)
    fila += 1
    _encabezado_tabla(hoja, fila, [(titulo.split(" por ")[-1], 1, 4), ("En plazo", 5, 5), ("Exigibles", 6, 6),
                                   ("Cumplimiento", 7, 8)])
    primera = fila + 1
    for i, dato in enumerate(filas):
        fila += 1
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        hoja.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=8)
        valores = {1: dato.etiqueta, 5: dato.en_plazo, 6: dato.exigibles,
                   7: (dato.porcentaje / 100) if dato.porcentaje is not None else None}
        for col in range(1, 9):
            celda = hoja.cell(row=fila, column=col)
            if col in valores:
                celda.value = valores[col]
            celda.font = Font(name=CUERPO, size=10)
            celda.border = BORDE_CELDA
            celda.alignment = Alignment(indent=1, vertical="center")
            if i % 2:
                celda.fill = relleno(CEBRA)
        hoja.cell(row=fila, column=7).number_format = "0.0%"
        hoja.row_dimensions[fila].height = 18
    if filas:
        hoja.conditional_formatting.add(
            f"G{primera}:G{fila}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=AZUL, showValue=True))
    else:
        fila += 1
        hoja.cell(row=fila, column=1, value="Sin revisiones exigibles en el período.").font = Font(
            name=CUERPO, size=10, italic=True, color=GRIS_TEXTO)
    return fila + 2


def _hoja_resumen(libro, ind):
    hoja = libro.active
    hoja.title = "Resumen"
    hoja.sheet_properties.tabColor = AZUL
    for col in range(1, 9):
        hoja.column_dimensions[get_column_letter(col)].width = 16
    _cabecera(hoja, "Reporte de cumplimiento", ind.filtros.etiqueta, 8)
    hoja.row_dimensions[6].height = 10
    hoja.row_dimensions[7].height = 4
    hoja.row_dimensions[8].height = 20
    hoja.row_dimensions[9].height = 34
    hoja.row_dimensions[10].height = 18
    _tarjeta(hoja, 1, "Cumplimiento del período",
             (ind.porcentaje / 100) if ind.porcentaje is not None else "—",
             f"{ind.en_plazo} de {ind.exigibles} revisiones en plazo", AZUL, "0.0%")
    _tarjeta(hoja, 3, "Revisiones vencidas", len(ind.vencidas), ind.detalle_vencidas or "Ninguna", ROJO_FALLA)
    _tarjeta(hoja, 5, "Hallazgos abiertos", len(ind.hallazgos_abiertos),
             f"{ind.hallazgos_alta} de criticidad alta", AMBAR_OBS)
    _tarjeta(hoja, 7, "Elementos activos sin plan", ind.sin_plan, f"de {ind.total_activos} elementos activos", NEGRO)
    hoja.cell(row=7, column=3).fill = hoja.cell(row=7, column=4).fill = relleno(ROJO)
    hoja.cell(row=7, column=5).fill = hoja.cell(row=7, column=6).fill = relleno(AMBAR)

    fila = 12
    hoja.cell(row=fila, column=1, value=(f"Realizadas fuera de plazo: {ind.fuera_de_plazo}   ·   "
                                         f"Revisiones con fecha límite en el período: {ind.total_periodo}")).font = Font(
        name=CUERPO, size=9, color=GRIS_TEXTO)
    fila = _tabla_cumplimiento(hoja, fila + 2, "Cumplimiento por unidad", ind.por_unidad)
    fila = _tabla_cumplimiento(hoja, fila, "Cumplimiento por categoría", ind.por_categoria)
    nota = hoja.cell(row=fila, column=1, value=("Cumplimiento = revisiones realizadas en plazo / revisiones con fecha "
                                                "límite en el período (RF-09). Las que aún están dentro de plazo no se "
                                                "consideran."))
    nota.font = Font(name=CUERPO, size=8, italic=True, color=GRIS_TEXTO)
    _preparar_impresion(hoja)


# ---------------------------------------------------------- Hojas de datos ----

def _hoja_tabla(libro, titulo, etiqueta_filtros, columnas, filas, color_pestana, estilo_celda=None):
    """columnas: lista de (encabezado, ancho, formato_numero)."""
    hoja = libro.create_sheet(titulo)
    hoja.sheet_properties.tabColor = color_pestana
    n = len(columnas)
    for i, (_, ancho, _) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    _cabecera(hoja, titulo, etiqueta_filtros, n)
    fila_enc = 7
    _encabezado_tabla(hoja, fila_enc, [(texto, i, i) for i, (texto, _, _) in enumerate(columnas, start=1)])
    fila = fila_enc
    for indice, valores in enumerate(filas):
        fila += 1
        hoja.row_dimensions[fila].height = 18
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.font = Font(name=CUERPO, size=10)
            celda.border = BORDE_CELDA
            formato = columnas[col - 1][2]
            if formato:  # fechas centradas, para que no se peguen a la columna vecina
                celda.number_format = formato
                celda.alignment = Alignment(vertical="center", horizontal="center")
            else:
                celda.alignment = Alignment(vertical="center", indent=1,
                                            wrap_text=columnas[col - 1][0] == "Descripción")
            if indice % 2:
                celda.fill = relleno(CEBRA)
            if estilo_celda:
                estilo_celda(columnas[col - 1][0], celda)
    if not filas:
        fila += 1
        hoja.cell(row=fila, column=1, value="Sin registros para estos filtros.").font = Font(
            name=CUERPO, size=10, italic=True, color=GRIS_TEXTO)
    else:
        hoja.auto_filter.ref = f"A{fila_enc}:{get_column_letter(n)}{fila}"
    hoja.freeze_panes = f"A{fila_enc + 1}"
    _preparar_impresion(hoja, f"{fila_enc}:{fila_enc}")
    return hoja


COLORES_RESULTADO = {"Conforme": VERDE_CONFORME, "Observación": AMBAR_OBS, "Falla": ROJO_FALLA}
COLORES_ESTADO = {"Vencida": ROJO_FALLA, "Realizada": VERDE_CONFORME, "No realizada": GRIS_TEXTO}
RELLENO_CRITICIDAD = {"Alta": ROJO_FALLA, "Media": AMBAR, "Baja": "6B7178"}


def _estilo_revisiones(columna, celda):
    if columna == "Resultado" and celda.value in COLORES_RESULTADO:
        celda.font = Font(name=CUERPO, size=10, bold=True, color=COLORES_RESULTADO[celda.value])
    elif columna == "Estado" and celda.value in COLORES_ESTADO:
        celda.font = Font(name=CUERPO, size=10, bold=celda.value == "Vencida", color=COLORES_ESTADO[celda.value])
    elif columna == "En plazo" and celda.value:
        celda.font = Font(name=CUERPO, size=10, bold=True,
                          color=VERDE_CONFORME if celda.value == "Sí" else ROJO_FALLA)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    elif columna == "Código":
        celda.font = Font(name="Consolas", size=10, bold=True)


def _estilo_hallazgos(columna, celda):
    if columna == "Criticidad" and celda.value in RELLENO_CRITICIDAD:
        celda.fill = relleno(RELLENO_CRITICIDAD[celda.value])
        celda.font = Font(name=TITULOS, size=10, bold=True, color="FFFFFF")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    elif columna == "Código":
        celda.font = Font(name="Consolas", size=10, bold=True)


def exportar_xlsx(indicadores, revisiones):
    libro = Workbook()
    libro.properties.title = f"Reporte de cumplimiento SIREMA · {indicadores.filtros.periodo}"
    libro.properties.creator = "SIREMA · Octava Compañía de Bomberos de Talcahuano"
    _hoja_resumen(libro, indicadores)

    filas_revisiones = []
    for r in revisiones:
        material = r.material
        en_plazo = r.en_plazo
        filas_revisiones.append([
            material.codigo_institucional or "s/c", material.nombre, material.categoria.nombre, material.ubicacion,
            str(r.plan), r.usuario_responsable.nombre_completo, _fecha(r.fecha_programada), _fecha(r.fecha_limite),
            r.get_estado_display(), r.get_resultado_global_display() if r.resultado_global else "",
            r.usuario_ejecutor.nombre_completo if r.usuario_ejecutor else "", _fecha(r.fecha_hora_termino),
            "" if en_plazo is None else ("Sí" if en_plazo else "No"),
        ])
    _hoja_tabla(libro, "Revisiones", indicadores.filtros.etiqueta, [
        ("Código", 11, None), ("Elemento", 30, None), ("Categoría", 14, None), ("Ubicación", 22, None),
        ("Plan", 18, None), ("Responsable", 22, None), ("Programada", 14, "dd-mm-yyyy"),
        ("Límite", 13, "dd-mm-yyyy"), ("Estado", 14, None), ("Resultado", 14, None), ("Ejecutada por", 22, None),
        ("Fecha de ejecución", 21, "dd-mm-yyyy hh:mm"), ("En plazo", 11, None),
    ], filas_revisiones, NEGRO, _estilo_revisiones)

    filas_hallazgos = []
    for h in indicadores.hallazgos_abiertos:
        material = h.revision.material
        filas_hallazgos.append([material.codigo_institucional or "s/c", material.nombre, material.ubicacion,
                                h.descripcion, h.get_criticidad_display(), h.get_estado_display(),
                                _fecha(h.fecha_apertura)])
    _hoja_tabla(libro, "Hallazgos abiertos", indicadores.filtros.etiqueta, [
        ("Código", 11, None), ("Elemento", 26, None), ("Ubicación", 22, None), ("Descripción", 60, None),
        ("Criticidad", 13, None), ("Estado", 15, None), ("Apertura", 19, "dd-mm-yyyy hh:mm"),
    ], filas_hallazgos, ROJO, _estilo_hallazgos)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
