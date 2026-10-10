"""
Carga masiva de material desde CSV o XLSX (CU-B1, curso alternativo 3.a).

Cada fila se valida con las mismas reglas del formulario. Las válidas se
registran (cada una en su propia transacción) y las rechazadas se informan
con su número de fila y el motivo.
"""
import csv
import io
import unicodedata
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError

from nucleo.enums import TipoMaterial

from . import services
from .models import Cajonera, Categoria

COLUMNAS = ["codigo_institucional", "nombre", "categoria", "unidad", "cajonera", "tipo", "cantidad",
            "marca", "modelo", "numero_serie", "observaciones"]
OBLIGATORIAS = ["nombre", "categoria", "unidad", "cajonera", "tipo"]
SINONIMOS = {"codigo": "codigo_institucional", "serie": "numero_serie", "n_serie": "numero_serie",
             "numero_de_serie": "numero_serie", "ubicacion": "cajonera", "carro": "unidad"}


@dataclass
class ResultadoImportacion:
    registrados: list = field(default_factory=list)
    rechazados: list = field(default_factory=list)

    @property
    def total(self):
        return len(self.registrados) + len(self.rechazados)


def _normalizar_encabezado(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.strip().lower().replace(" ", "_").replace(".", "")
    return SINONIMOS.get(texto, texto)


def _leer_csv(contenido):
    texto = contenido.decode("utf-8-sig", errors="replace")
    try:
        dialecto = csv.Sniffer().sniff(texto.splitlines()[0] if texto else "", delimiters=";,\t")
    except csv.Error:
        dialecto = csv.excel
    return list(csv.reader(io.StringIO(texto), dialecto))


def _leer_xlsx(contenido):
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.active
    return [["" if c is None else str(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]


def leer_filas(archivo):
    nombre = (getattr(archivo, "name", "") or "").lower()
    contenido = archivo.read()
    if nombre.endswith(".xlsx"):
        try:
            return _leer_xlsx(contenido)
        except Exception as exc:  # openpyxl lanza varios tipos según el daño del archivo
            raise ValidationError("No se pudo leer la planilla XLSX.") from exc
    if nombre.endswith(".csv"):
        return _leer_csv(contenido)
    raise ValidationError("Formato no admitido: adjunte una planilla CSV o XLSX.")


def _fila_a_datos(fila):
    errores = []
    for columna in OBLIGATORIAS:
        if not fila.get(columna):
            errores.append(f"falta «{columna}»")
    if errores:
        return None, errores

    categoria = Categoria.objects.filter(nombre__iexact=fila["categoria"]).first()
    if categoria is None:
        errores.append(f"la categoría «{fila['categoria']}» no existe")
    cajonera = Cajonera.objects.select_related("unidad").filter(
        unidad__codigo__iexact=fila["unidad"], nombre__iexact=fila["cajonera"]).first()
    if cajonera is None:
        errores.append(f"la cajonera «{fila['unidad']} / {fila['cajonera']}» no existe")
    tipo = fila["tipo"].strip().lower()
    if tipo not in TipoMaterial.values:
        errores.append("el tipo debe ser «unitario» o «granel»")
    cantidad = 1
    if fila.get("cantidad"):
        try:
            cantidad = int(float(fila["cantidad"].replace(",", ".")))
            if cantidad < 1:
                raise ValueError
        except ValueError:
            errores.append("la cantidad debe ser un entero mayor o igual a 1")
    if errores:
        return None, errores

    datos = {
        "codigo_institucional": fila.get("codigo_institucional") or None,
        "nombre": fila["nombre"][:100],
        "categoria": categoria,
        "cajonera": cajonera,
        "tipo": tipo,
        "cantidad": cantidad,
        "marca": (fila.get("marca") or None),
        "modelo": (fila.get("modelo") or None),
        "numero_serie": (fila.get("numero_serie") or None),
        "observaciones": (fila.get("observaciones") or None),
    }
    return datos, []


def importar_planilla(archivo, usuario):
    filas = leer_filas(archivo)
    if not filas:
        raise ValidationError("La planilla está vacía.")
    encabezados = [_normalizar_encabezado(c) for c in filas[0]]
    faltantes = [c for c in OBLIGATORIAS if c not in encabezados]
    if faltantes:
        raise ValidationError("Faltan columnas obligatorias: " + ", ".join(faltantes) + ".")

    resultado = ResultadoImportacion()
    for numero, valores in enumerate(filas[1:], start=2):
        fila = {encabezados[i]: str(v).strip() for i, v in enumerate(valores) if i < len(encabezados)}
        if not any(fila.values()):
            continue
        datos, errores = _fila_a_datos(fila)
        if errores:
            resultado.rechazados.append({"fila": numero, "nombre": fila.get("nombre", ""), "motivos": errores})
            continue
        try:
            material = services.registrar_material(datos, usuario)
        except ValidationError as error:
            resultado.rechazados.append({"fila": numero, "nombre": fila.get("nombre", ""),
                                         "motivos": error.messages})
            continue
        resultado.registrados.append(material)
    return resultado


def plantilla_csv():
    ejemplo = ["8C-0999", "Motobomba portátil", "MOTORES", "B-8", "Cabina trasera", "unitario", "1",
               "Honda", "WX15", "", "Ejemplo: borre esta fila"]
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow(COLUMNAS)
    escritor.writerow(ejemplo)
    return "﻿" + salida.getvalue()
